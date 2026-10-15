import io

import pytest

pptx = pytest.importorskip("pptx")
from pptx import Presentation
from pptx.util import Inches, Pt

import deck_reader as dr

# a 1x1 white PNG, enough for python-pptx to place a picture
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c63f8ffffff7f0009fb03fd2a86e38a"
    "0000000049454e44ae426082"
)


def _box(slide, text, size=None, top=0):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5 + top), Inches(8), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    if size is not None:
        run.font.size = Pt(size)
    return box


def _deck(tmp_path, width_in=13.333):
    prs = Presentation()
    prs.slide_width = Inches(width_in)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    cover = prs.slides.add_slide(blank)
    _box(cover, "Market entry plan", size=40)

    body = prs.slides.add_slide(blank)
    _box(body, "Revenue grows 20% a year", size=24)
    _box(body, "source: made-up sample data", size=6, top=2)
    body.shapes.add_picture(io.BytesIO(PNG), Inches(9), Inches(1), Inches(2), Inches(2))

    table_slide = prs.slides.add_slide(blank)
    table = table_slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2)).table
    for i, cell in enumerate(c for row in table.rows for c in row.cells):
        cell.text = f"cell {i}"
        cell.text_frame.paragraphs[0].runs[0].font.size = Pt(12)

    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_reads_words_images_and_sizes_per_slide(tmp_path):
    slides = dr.read_pptx(_deck(tmp_path))
    assert [s.words for s in slides] == [3, 9, 8]
    assert [s.images for s in slides] == [0, 1, 0]
    assert slides[0].body_pt == pytest.approx(40, abs=0.1)
    assert slides[1].smallest_pt == pytest.approx(6, abs=0.1)
    assert slides[2].body_pt == pytest.approx(12, abs=0.1)


def test_sizes_are_scaled_from_a_narrower_slide(tmp_path):
    # 40pt on a 10in-wide slide is 53.3pt on a 13.33in-wide one
    slides = dr.read_pptx(_deck(tmp_path, width_in=10))
    assert slides[0].body_pt == pytest.approx(53.3, abs=0.1)


def test_run_without_a_size_falls_back_to_the_default(tmp_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    _box(prs.slides.add_slide(prs.slide_layouts[6]), "no size set here")
    path = tmp_path / "plain.pptx"
    prs.save(path)
    (slide,) = dr.read_pptx(path, default_pt=14)
    assert slide.words == 4 and slide.body_pt == pytest.approx(14, abs=0.1)


def test_small_type_in_a_pptx_is_flagged(tmp_path):
    slides = dr.read_pptx(_deck(tmp_path))
    assert dr.slide_flags(slides[1]) == ["type at 6.0pt, under the 8pt floor"]


def test_cli_reads_pptx_by_extension(tmp_path, capsys):
    assert dr.main([str(_deck(tmp_path)), "--format", "csv", "--strict"]) == 1
    rows = capsys.readouterr().out.splitlines()
    assert rows[0].startswith("deck,slide")
    assert rows[2].startswith("deck.pptx,2,9,1,")


def test_cli_refuses_to_render_a_pptx(tmp_path):
    with pytest.raises(SystemExit):
        dr.main([str(_deck(tmp_path)), str(tmp_path / "out")])
