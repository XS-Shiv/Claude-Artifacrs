import json

import pytest

import deck_reader as dr
from deck_reader import SlideStats

DECK = [SlideStats(1, 4, 0, 40.0, 40.0), SlideStats(2, 120, 0, 12.0, 10.0), SlideStats(3, 90, 1, 12.0, 12.0)]


def test_json_without_titles_is_unchanged():
    doc = json.loads(dr.format_json("a.pdf", DECK))
    assert "topic_only_titles" not in doc
    assert "title" not in doc["per_slide"][0]


def test_titles_are_matched_to_slides_by_number():
    titles = [(1, "Market entry plan"), (2, "Revenue grows 20% a year")]
    doc = json.loads(dr.format_json("a.pdf", DECK, titles=titles))
    assert [(r["slide"], r["title"], r["topic_only"]) for r in doc["per_slide"]] == [
        (1, "Market entry plan", True),
        (2, "Revenue grows 20% a year", False),
        (3, None, None),
    ]
    assert doc["topic_only_titles"] == [1]


def test_summary_keys_come_before_the_per_slide_list():
    doc = json.loads(dr.format_json("a.pdf", DECK, titles=[]))
    assert list(doc)[-2:] == ["topic_only_titles", "per_slide"]
    assert doc["topic_only_titles"] == []


def test_topic_only_titles_do_not_count_as_readability_flags():
    doc = json.loads(dr.format_json("a.pdf", DECK, titles=[(2, "Competitor landscape")]))
    assert doc["topic_only_titles"] == [2]
    assert doc["flagged_slides"] == []


def test_cli_json_carries_pptx_titles(tmp_path, capsys):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for title in ["Chennai market analysis", "Chennai resists quick commerce but rewards first movers"]:
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = title
    path = tmp_path / "titled.pptx"
    prs.save(path)

    assert dr.main([str(path), "--format", "json"]) == 0
    doc = json.loads(capsys.readouterr().out)
    assert doc["topic_only_titles"] == [1]
    assert [r["topic_only"] for r in doc["per_slide"]] == [True, False]
    assert doc["per_slide"][0]["title"] == "Chennai market analysis"
