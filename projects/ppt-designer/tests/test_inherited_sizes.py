import pytest

pytest.importorskip("pptx")
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

import deck_reader as dr

A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _set_list_style(shape, level, pt):
    """Give a placeholder an a:lstStyle size for one paragraph level."""
    body = shape._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}txBody")
    lst = body.find(f"{{{A}}}lstStyle")
    if lst is None:
        lst = etree.SubElement(body, f"{{{A}}}lstStyle")
        body.remove(lst)
        body.insert(1, lst)
    lvl = etree.SubElement(lst, f"{{{A}}}lvl{level + 1}pPr")
    etree.SubElement(lvl, f"{{{A}}}defRPr", sz=str(int(pt * 100)))


def _title_and_content(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    frame = slide.placeholders[1].text_frame
    frame.text = bullets[0][0]
    frame.paragraphs[0].level = bullets[0][1]
    for text, level in bullets[1:]:
        para = frame.add_paragraph()
        para.text = text
        para.level = level
    return slide


def test_master_text_styles_give_title_and_body_levels():
    prs = Presentation()
    slide = _title_and_content(prs, "Sales rose", [("Top level point", 0), ("Second level", 1)])
    assert dr.inherited_pt(slide.shapes.title) == 44.0
    body = slide.placeholders[1]
    assert dr.pptx_word_sizes(body.text_frame, shape=body) == [32.0] * 3 + [28.0] * 2


def test_layout_placeholder_wins_over_master():
    prs = Presentation()
    _set_list_style(prs.slide_layouts[1].placeholders[1], 0, 20)
    slide = _title_and_content(prs, "Sales rose", [("Point", 0), ("Sub point", 1)])
    body = slide.placeholders[1]
    # level 0 comes from the layout, level 1 still from the master body style
    assert dr.pptx_word_sizes(body.text_frame, shape=body) == [20.0, 28.0, 28.0]


def test_master_placeholder_wins_over_text_styles():
    prs = Presentation()
    master_title = next(p for p in prs.slide_master.placeholders if p.placeholder_format.idx == 0)
    _set_list_style(master_title, 0, 36)
    slide = _title_and_content(prs, "Sales rose", [("Point", 0)])
    assert dr.inherited_pt(slide.shapes.title) == 36.0


def test_slide_list_style_and_explicit_sizes_win():
    prs = Presentation()
    slide = _title_and_content(prs, "Sales rose", [("Point", 0)])
    body = slide.placeholders[1]
    _set_list_style(body, 0, 14)
    assert dr.inherited_pt(body) == 14.0
    body.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    assert dr.pptx_word_sizes(body.text_frame, shape=body) == [10.0]


def test_text_boxes_keep_the_default():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "two words"
    assert dr.inherited_pt(box) is None
    assert dr.pptx_word_sizes(box.text_frame, shape=box) == [18.0, 18.0]


def test_read_pptx_uses_inherited_sizes(tmp_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    _title_and_content(prs, "Sales rose", [("Small print here", 1)])
    path = tmp_path / "deck.pptx"
    prs.save(path)
    (stats,) = dr.read_pptx(path)
    assert stats.words == 5
    assert stats.smallest_pt == pytest.approx(28.0, abs=0.1)
    assert max(dr.pptx_word_sizes(prs.slides[0].shapes.title.text_frame, shape=prs.slides[0].shapes.title)) == 44.0
