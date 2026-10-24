"""Measure a slide deck, as PDF or .pptx: words, images and type sizes on every slide.

Type sizes are reported as they would appear on a standard 13.33in-wide slide,
so decks exported at different page sizes can be compared with each other.

Slides that break a readability limit are listed under the table: type smaller than
8pt, or more than 450 words. Both limits come from the skill in SKILL.md.

Usage:
    python deck_reader.py deck.pdf            print the table and any slides over the limits
    python deck_reader.py deck.pdf out_dir    also render every slide into out_dir as JPEG
    python deck_reader.py deck.pdf --strict   exit with status 1 when a slide is over a limit
    python deck_reader.py deck.pdf --format json   (or csv) print machine-readable output;
                                              JSON also carries each slide's title check
    python deck_reader.py deck.pptx           read a PowerPoint file directly
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import shutil
import statistics
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

SLIDE_WIDTH_PT = 960.0  # a 13.33in-wide slide at 72 points per inch
CONTENT_MIN_WORDS = 60  # fewer words than this is a cover, divider or thank-you slide
MIN_TYPE_PT = 8.0  # the skill's floor: smaller type cannot be read on a slide
MAX_WORDS = 450  # the densest content slides in the skill's reference decks
DEFAULT_PPTX_PT = 18.0  # PowerPoint's default body size, used when a run sets no size
EMU_PER_PT = 12700


@dataclass(frozen=True)
class SlideStats:
    number: int
    words: int
    images: int
    body_pt: float  # median type size on the slide
    smallest_pt: float


def scale_to_slide(size_pt: float, page_width_pt: float) -> float:
    """Convert a type size on a page of any width to its size on a standard slide."""
    if page_width_pt <= 0:
        raise ValueError("page width must be positive")
    return size_pt * SLIDE_WIDTH_PT / page_width_pt


def slide_stats(
    number: int, word_sizes: Iterable[float], page_width_pt: float, images: int = 0
) -> SlideStats:
    """Summarise one slide from the type size of each of its words."""
    sizes = [scale_to_slide(size, page_width_pt) for size in word_sizes]
    if not sizes:
        return SlideStats(number, 0, images, 0.0, 0.0)
    return SlideStats(number, len(sizes), images, statistics.median(sizes), min(sizes))


def content_slides(
    slides: Sequence[SlideStats], min_words: int = CONTENT_MIN_WORDS
) -> list[SlideStats]:
    """Drop covers, dividers and thank-you slides, which carry too few words to count."""
    return [slide for slide in slides if slide.words >= min_words]


def median_words(slides: Sequence[SlideStats]) -> float | None:
    """Median words per content slide, or None when the deck has no content slides."""
    counts = [slide.words for slide in content_slides(slides)]
    return statistics.median(counts) if counts else None


def read_pdf(path: str | Path) -> list[SlideStats]:
    """Read every page of a PDF deck into SlideStats."""
    import pdfplumber  # imported here so the functions above work without it installed

    slides = []
    with pdfplumber.open(path) as pdf:
        for number, page in enumerate(pdf.pages, 1):
            words = page.extract_words(extra_attrs=["size"])
            sizes = [word["size"] for word in words]
            slides.append(slide_stats(number, sizes, float(page.width), len(page.images)))
    return slides


def _pptx_shapes(shapes):
    """Every shape on a slide, looking inside groups."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pptx_shapes(shape.shapes)
        else:
            yield shape


def _text_frames(shape):
    """The text frames of a shape: its own, or one per cell of a table."""
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        yield shape.text_frame
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text_frame


_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def _list_style_pt(element, level: int) -> float | None:
    """The size an a:lstStyle (or a master text style) sets for a paragraph level."""
    if element is None:
        return None
    lvl = element.find(f"{_A}lvl{level + 1}pPr")
    rpr = lvl.find(f"{_A}defRPr") if lvl is not None else None
    if rpr is not None and rpr.get("sz"):
        return int(rpr.get("sz")) / 100
    return None


def _shape_list_style(shape):
    body = shape._element.find(f".//{_P}txBody")
    return body.find(f"{_A}lstStyle") if body is not None else None


def inherited_pt(shape, level: int = 0) -> float | None:
    """The type size a placeholder inherits when its runs and paragraphs set none.

    Follows PowerPoint's order: the shape's own list style, the matching
    placeholder on the slide layout (same idx), the master placeholder of the
    same kind, then the master's title, body or other text style. Returns None
    for shapes that are not placeholders, and when nothing in the chain sets a size.
    """
    from pptx.enum.shapes import PP_PLACEHOLDER as PH

    if not getattr(shape, "is_placeholder", False):
        return None
    titles = {PH.TITLE, PH.CENTER_TITLE, PH.VERTICAL_TITLE}
    footers = {PH.DATE, PH.FOOTER, PH.SLIDE_NUMBER}
    chain = [_shape_list_style(shape)]
    kind = shape.placeholder_format.type
    layout = getattr(shape.part, "slide_layout", None)
    if layout is not None:
        idx = shape.placeholder_format.idx
        match = next((p for p in layout.placeholders if p.placeholder_format.idx == idx), None)
        if match is not None:
            chain.append(_shape_list_style(match))
            kind = match.placeholder_format.type
        master = layout.slide_master
        want = PH.TITLE if kind in titles else kind if kind in footers else PH.BODY
        mph = next((p for p in master.placeholders if p.placeholder_format.type == want), None)
        if mph is not None:
            chain.append(_shape_list_style(mph))
        styles = master._element.find(f"{_P}txStyles")
        if styles is not None:
            name = "titleStyle" if kind in titles else "otherStyle" if kind in footers else "bodyStyle"
            chain.append(styles.find(f"{_P}{name}"))
    for element in chain:
        pt = _list_style_pt(element, level)
        if pt is not None:
            return pt
    return None


def pptx_word_sizes(text_frame, default_pt: float = DEFAULT_PPTX_PT, shape=None) -> list[float]:
    """The type size of each word in a text frame, in points.

    A run's own size wins, then its paragraph's, then the size the shape
    inherits from the slide layout and master when it is a placeholder;
    anything left unset gets default_pt.
    """
    sizes = []
    for paragraph in text_frame.paragraphs:
        para_size = paragraph.font.size
        fallback = None
        for run in paragraph.runs:
            size = run.font.size or para_size
            if size is not None:
                pt = size.pt
            else:
                if fallback is None:
                    fallback = inherited_pt(shape, paragraph.level) if shape is not None else None
                    fallback = fallback if fallback is not None else default_pt
                pt = fallback
            sizes.extend([pt] * len(run.text.split()))
    return sizes


def read_pptx(path: str | Path, default_pt: float = DEFAULT_PPTX_PT) -> list[SlideStats]:
    """Read every slide of a .pptx deck into SlideStats, without converting it to PDF."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    width_pt = prs.slide_width / EMU_PER_PT
    slides = []
    for number, slide in enumerate(prs.slides, 1):
        sizes: list[float] = []
        images = 0
        for shape in _pptx_shapes(slide.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                images += 1
            for frame in _text_frames(shape):
                sizes.extend(pptx_word_sizes(frame, default_pt, shape))
        slides.append(slide_stats(number, sizes, width_pt, images))
    return slides


def read_deck(path: str | Path) -> list[SlideStats]:
    """Read a deck by its extension: .pptx with python-pptx, anything else as PDF."""
    return read_pptx(path) if Path(path).suffix.lower() == ".pptx" else read_pdf(path)


def render_slides(path: str | Path, out_dir: str | Path, width_px: int = 1280) -> None:
    """Write every slide into out_dir as s-<n>.jpg, using poppler's pdftoppm."""
    if shutil.which("pdftoppm") is None:
        raise RuntimeError("pdftoppm was not found; install poppler to render slides")
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["pdftoppm", "-scale-to-x", str(width_px), "-scale-to-y", "-1", "-jpeg",
         str(path), str(out / "s")],
        check=True,
    )


def _count(n: int, noun: str) -> str:
    return f"{n} {noun}" if n == 1 else f"{n} {noun}s"


def format_report(slides: Sequence[SlideStats]) -> str:
    """One summary line, then one line per slide."""
    median = median_words(slides)
    summary = f"{_count(len(slides), 'slide')}, {_count(len(content_slides(slides)), 'content slide')}"
    if median is not None:
        summary += f", median {median:.0f} words per content slide"
    lines = [summary]
    for s in slides:
        lines.append(
            f"{s.number:>2} | {s.words:>4} words | {s.images:>2} images"
            f" | body {s.body_pt:4.1f}pt | smallest {s.smallest_pt:4.1f}pt"
        )
    return "\n".join(lines)


def slide_flags(
    slide: SlideStats, min_type_pt: float = MIN_TYPE_PT, max_words: int = MAX_WORDS
) -> list[str]:
    """The readability limits one slide breaks, as short phrases. Empty when it breaks none."""
    flags = []
    # compare the size as the report prints it, so 7.96pt shows and counts as 8.0pt
    if slide.words and round(slide.smallest_pt, 1) < min_type_pt:
        flags.append(f"type at {slide.smallest_pt:.1f}pt, under the {min_type_pt:g}pt floor")
    if slide.words > max_words:
        flags.append(f"{slide.words} words, over the {max_words}-word limit")
    return flags


def format_flags(
    slides: Sequence[SlideStats], min_type_pt: float = MIN_TYPE_PT, max_words: int = MAX_WORDS
) -> str:
    """One line per broken limit under a heading, or an empty string when nothing is flagged."""
    lines = [
        f"  slide {slide.number}: {flag}"
        for slide in slides
        for flag in slide_flags(slide, min_type_pt, max_words)
    ]
    return "\n".join(["Over the limits:", *lines]) if lines else ""


CSV_FIELDS = ["deck", "slide", "words", "images", "body_pt", "smallest_pt", "content", "flags"]


def to_records(
    slides: Sequence[SlideStats], min_type_pt: float = MIN_TYPE_PT, max_words: int = MAX_WORDS
) -> list[dict]:
    """One dict per slide, with sizes rounded as the text report shows them."""
    return [
        {
            "slide": s.number,
            "words": s.words,
            "images": s.images,
            "body_pt": round(s.body_pt, 1),
            "smallest_pt": round(s.smallest_pt, 1),
            "content": s.words >= CONTENT_MIN_WORDS,
            "flags": slide_flags(s, min_type_pt, max_words),
        }
        for s in slides
    ]


def add_titles(records: list[dict], titles: Sequence[tuple[int, str]]) -> list[int]:
    """Put each slide's title and its topic-only verdict into its record.

    Slides with no title found get None for both. Returns the topic-only slide numbers.
    """
    from titles import is_topic_only

    by_slide = dict(titles)
    for record in records:
        title = by_slide.get(record["slide"])
        record["title"] = title
        record["topic_only"] = None if title is None else is_topic_only(title)
    return [r["slide"] for r in records if r["topic_only"]]


def format_json(
    deck: str, slides: Sequence[SlideStats],
    min_type_pt: float = MIN_TYPE_PT, max_words: int = MAX_WORDS,
    titles: Sequence[tuple[int, str]] | None = None,
) -> str:
    """The deck's summary, limits and per-slide measurements as one JSON document.

    With titles, as (slide number, title) pairs from titles.read_titles, each slide also
    carries its title and topic-only verdict, and the summary lists the topic-only slides.
    """
    records = to_records(slides, min_type_pt, max_words)
    topic_only = add_titles(records, titles) if titles is not None else None
    doc = {
        "deck": deck,
        "slides": len(slides),
        "content_slides": len(content_slides(slides)),
        "median_words": median_words(slides),
        "limits": {"min_type_pt": min_type_pt, "max_words": max_words},
        "flagged_slides": [r["slide"] for r in records if r["flags"]],
    }
    if topic_only is not None:
        doc["topic_only_titles"] = topic_only
    doc["per_slide"] = records
    return json.dumps(doc, indent=2)


def format_csv(
    deck: str, slides: Sequence[SlideStats], min_type_pt: float = MIN_TYPE_PT,
    max_words: int = MAX_WORDS, header: bool = True,
) -> str:
    """One row per slide, with the deck name in every row so several decks can be stacked."""
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=CSV_FIELDS, lineterminator="\n")
    if header:
        writer.writeheader()
    for record in to_records(slides, min_type_pt, max_words):
        writer.writerow({**record, "deck": deck, "flags": "; ".join(record["flags"])})
    return buf.getvalue().rstrip("\n")


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Measure a slide deck slide by slide.")
    parser.add_argument("deck", help="the deck as a PDF or .pptx")
    parser.add_argument("out_dir", nargs="?", help="folder to render the slides into as JPEG")
    parser.add_argument("--min-type", type=float, default=MIN_TYPE_PT, metavar="PT",
                        help="flag slides with type smaller than this (default: %(default)g)")
    parser.add_argument("--max-words", type=int, default=MAX_WORDS, metavar="N",
                        help="flag slides with more words than this (default: %(default)s)")
    parser.add_argument("--strict", action="store_true",
                        help="exit with status 1 when any slide is over a limit")
    parser.add_argument("--format", choices=["text", "json", "csv"], default="text",
                        help="output format (default: %(default)s)")
    parser.add_argument("--no-header", action="store_true",
                        help="with --format csv, leave out the header row, to append to a file")
    args = parser.parse_args(argv)
    if not Path(args.deck).is_file():
        parser.error(f"no such file: {args.deck}")
    if args.out_dir and Path(args.deck).suffix.lower() == ".pptx":
        parser.error("rendering slides needs a PDF; export the .pptx to PDF first")
    slides = read_deck(args.deck)
    if args.out_dir:
        render_slides(args.deck, args.out_dir)
    flags = format_flags(slides, args.min_type, args.max_words)
    name = Path(args.deck).name
    if args.format == "json":
        from titles import read_titles

        print(format_json(name, slides, args.min_type, args.max_words, read_titles(args.deck)))
    elif args.format == "csv":
        print(format_csv(name, slides, args.min_type, args.max_words, not args.no_header))
    else:
        print(format_report(slides))
        if flags:
            print(flags)
    return 1 if flags and args.strict else 0


if __name__ == "__main__":
    sys.exit(main())
