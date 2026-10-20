"""`check` re-opens a morphed deck and lists what an edit after `apply` broke."""
import sys

import pytest

pytest.importorskip("pptx")
from pptx import Presentation
from pptx.util import Inches

import morph
from test_morph import H, W, _deck, _run


def _morphed(tmp_path, monkeypatch):
    out = tmp_path / "deck-morph.pptx"
    assert _run(monkeypatch, "apply", _deck(tmp_path / "deck.pptx"), "-o", out, "--report", tmp_path / "r.json") == 0
    return out


def _problems(path):
    return morph.check_deck(morph.Pkg(path))


def _edit(path, slide, fn):
    """Change one slide's XML in place, the way an edit in PowerPoint would."""
    pkg = morph.Pkg(path)
    part = pkg.slides[slide - 1]
    fn(pkg.xml(part))
    pkg.dirty.add(part)
    pkg.save(path)


def _named(root, start):
    return [c for c in root.iter(morph.q("p:cNvPr")) if c.get("name", "").startswith(start)]


def test_fresh_apply_has_no_problems(tmp_path, monkeypatch, capsys):
    out = _morphed(tmp_path, monkeypatch)
    assert _problems(out) == (2, [])
    assert _run(monkeypatch, "check", out) == 0
    assert "2 Morph transitions  |  no problems found" in capsys.readouterr().out


def test_transition_changed_back_is_flagged(tmp_path, monkeypatch, capsys):
    out = _morphed(tmp_path, monkeypatch)
    _edit(out, 2, morph.set_transition)           # no duration: drops the Morph transition
    n, bad = _problems(out)
    assert n == 1
    assert bad == ["slide 2: shares 3 morph name(s) with slide 1 but its transition is not Morph"]
    assert _run(monkeypatch, "check", out) == 1
    assert "1 problem(s)" in capsys.readouterr().out


def test_duplicate_name_is_flagged(tmp_path, monkeypatch):
    out = _morphed(tmp_path, monkeypatch)

    def clash(root):
        title, card = [c for c in _named(root, "!!m") if not morph.is_ghost(c.getparent().getparent(), W, H)][:2]
        card.set("name", title.get("name"))
    _edit(out, 2, clash)
    assert any("is used by 2 sp shapes" in line and line.startswith("slide 2:") for line in _problems(out)[1])


def test_deleted_shape_leaves_orphan_parked_copies(tmp_path, monkeypatch):
    out = _morphed(tmp_path, monkeypatch)

    def delete_note(root):
        note = next(c for c in _named(root, "!!m") if "Only here" in "".join(c.getparent().getparent().itertext()))
        shape = note.getparent().getparent()
        shape.getparent().remove(shape)
    _edit(out, 2, delete_note)
    bad = _problems(out)[1]
    assert sorted(line.split(":")[0] for line in bad) == ["slide 1", "slide 3"]
    assert all("parked copy" in line and "no partner" in line for line in bad)


def test_kind_mismatch_is_flagged(tmp_path, monkeypatch):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    for kind in ("box", "line"):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        if kind == "box":
            shape = s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
            shape.text_frame.text = "Arrow"
        else:
            shape = s.shapes.add_connector(1, Inches(1), Inches(1), Inches(4), Inches(2))
        shape.name = "!!arrow"
    deck, out = tmp_path / "k.pptx", tmp_path / "k-morph.pptx"
    prs.save(deck)
    # with fade there are no parked copies, so the two hand-named shapes face each other directly
    assert _run(monkeypatch, "apply", deck, "-o", out, "--style", "fade", "--report", tmp_path / "r.json") == 0
    assert _problems(out)[1] == ["slides 1 -> 2: '!!arrow' is a sp on one and a cxnSp on the other, so it cannot morph"]
